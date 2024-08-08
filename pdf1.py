import os
import shutil
import uuid
import secrets
import io
import jwt
import bcrypt
from flask import Flask, render_template, request, jsonify, session
from flask_sqlalchemy import SQLAlchemy
from PyPDF2 import PdfReader
from pptx import Presentation
from docx import Document
from dotenv import load_dotenv
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_community.vectorstores import FAISS
from langchain.prompts import PromptTemplate
from langchain_openai import OpenAIEmbeddings, ChatOpenAI
from langchain.chains.question_answering import load_qa_chain
import logging
from logging.handlers import RotatingFileHandler
from langchain_community.callbacks import get_openai_callback
from datetime import datetime, timedelta
from flask_cors import CORS
from userdb import db,UserInfo,UserActivities,generate_token,log_activities,signin,register,forgot_password

# Load environment variables
load_dotenv()

# Initialize Flask app
app = Flask(__name__)
CORS(app)

# SQLAlchemy configuration
app.config['SQLALCHEMY_DATABASE_URI'] = 'mysql+pymysql://root:1234@localhost:3306/user_credentials'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['SECRET_KEY'] = 'tucEDtE44BbQLv7tXCivZkn1DbmKGsYn'
db.init_app(app)  # Initialize the db with the Flask app

# Configure logging
log_filename = 'pdf_app.log'
handler = RotatingFileHandler(log_filename, maxBytes=1000000, backupCount=5)
formatter = logging.Formatter('%(asctime)s - %(name)s - %(levelname)s - %(message)s')
handler.setFormatter(formatter)
logger = logging.getLogger('pdf_app')
logger.setLevel(logging.DEBUG)
logger.addHandler(handler)

# Initialize the database
with app.app_context():
    db.create_all()

def auth():
    auth_header = request.headers.get('Authorization')
    
    if auth_header:
        try:
            token = auth_header.split(" ")[1]
            decoded = jwt.decode(token, app.config['SECRET_KEY'], algorithms=['HS256'])
            
            log_activities(
                user_id=decoded.get('user_id'),
                email=decoded.get('email'),
                activity_type="Token Validation",
                activity_content="Token validated successfully",
                response_content="Token is valid",
                error_message=None,
                app_name=None
            )
            
            return jsonify({"status": "Success", "data": decoded}), 200
        
        except jwt.ExpiredSignatureError:
            log_activities(
                user_id=None,
                email=None,
                activity_type="Token Validation",
                activity_content="Token expired",
                response_content="Token is expired",
                error_message="Token expired"
            )
            return jsonify({"status": "Failure", "message": "Token expired"}), 401
        
        except jwt.InvalidTokenError:
            log_activities(
                user_id=None,
                email=None,
                activity_type="Token Validation",
                activity_content="Invalid token",
                response_content="Token is invalid",
                error_message="Invalid token"
            )
            return jsonify({"status": "Failure", "message": "Invalid token"}), 401
    
    else:
        log_activities(
            user_id=None,
            email=None,
            activity_type="Token Validation",
            activity_content="Token is missing",
            response_content="No token provided",
            error_message="Token is missing"
        )
        return jsonify({"status": "Failure", "message": "Token is missing"}), 401


api_key2 = "sk-proj-KHzaTujZGduAhDVL2GyfT3BlbkFJeLWhvD8web0AbEnkxcJc"
chat_history = []


def extract_text_from_pdf(pdf_content):
    try:
        pdf_stream = io.BytesIO(pdf_content)
        pdf_reader = PdfReader(pdf_stream)
        text = ""
        for page_number, page in enumerate(pdf_reader.pages, start=1):
            text += f"Page {page_number}:\n"
            text += page.extract_text()
        return text
    except Exception as e:
        logger.error(f"Error extracting text from PDF: {str(e)}")
        return f"Error extracting text from PDF: {str(e)}"

def extract_text_from_docx(docx_file):
    try:
        doc = Document(docx_file)
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        return text
    except Exception as e:
        logger.error(f"Error extracting text from DOCX: {str(e)}")
        return f"Error extracting text from DOCX: {str(e)}"

def extract_text_from_pptx(pptx_file):
    try:
        prs = Presentation(pptx_file)
        text = ""
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
        return text
    except Exception as e:
        logger.error(f"Error extracting text from PPTX: {str(e)}")
        return f"Error extracting text from PPTX: {str(e)}"

def get_text_chunks(text):
    try:
        text_splitter = RecursiveCharacterTextSplitter(chunk_size=10000, chunk_overlap=1000)
        chunks = text_splitter.split_text(text)
        return chunks
    except Exception as e:
        logger.error(f"Error splitting text into chunks: {str(e)}")
        return f"Error splitting text into chunks: {str(e)}"

def get_vector_store(text_chunks, vector_store_path):
    try:
        embeddings = OpenAIEmbeddings(openai_api_key=api_key2)
        vector_store = FAISS.from_texts(text_chunks, embeddings)
        vector_store.save_local(vector_store_path)
        logger.info(f"Faiss index for user saved successfully at: {vector_store_path}")
    except Exception as e:
        logger.error(f"Error creating vector store: {str(e)}")
        return f"Error creating vector store: {str(e)}"

def delete_faiss_index(user_id):
    try:
        user_vector_store_path = f"persistent_storage/faiss_index_{user_id}"
        if os.path.exists(user_vector_store_path):
            shutil.rmtree(user_vector_store_path)
            logger.info(f"Successfully cleared FAISS index for user {user_id}")
        else:
            logger.warning(f"FAISS index directory not found for user {user_id}")
    except Exception as e:
        logger.error(f"Error deleting FAISS index: {str(e)}")

def get_conversation_chain():
    try:
        prompt_template = """
        Answer the question based on the content of the uploaded files. 
        If the answer is not available in the files, respond with "Answer is not available in the uploaded files." 
        Maintain a formal tone with proper grammar and punctuation. 
        Summarize visual content if the question refers to diagrams or images. 
        If the question is unclear, ask for clarification. 
        Provide answers point by point, avoid unnecessary details. 
        Keep your response within two lines.
        If the user has uploaded files in unsupported formats, respond with "Please upload files in supported formats: PDF, DOC, DOCX, PPTX, PPT, or TXT."
        
        {context}

        Question:
        {question}

        Answer:
        """
        model = ChatOpenAI(model="gpt-3.5-turbo", temperature=0, openai_api_key=api_key2)
        prompt = PromptTemplate(template=prompt_template, input_variables=["context", "question"])
        chain = load_qa_chain(model, chain_type="stuff", prompt=prompt)
        return chain
    except Exception as e:
        logger.error(f"Error getting conversation chain: {str(e)}")
        return f"Error getting conversation chain: {str(e)}"

def get_question_generation_chain():
    try:
        prompt_template = """
        Based on the content of the uploaded files or folder, generate the top 5 questions that a user might ask. Don't extract the questions in vector database. Generate the questions in uploaded files only. Don't provide wrong information. Generate the questions from all the files within the uploaded folder.
        {context}

        Questions:
        """
        model = ChatOpenAI(model="gpt-3.5-turbo", temperature=0, openai_api_key=api_key2)
        prompt = PromptTemplate(template=prompt_template, input_variables=["context"])
        chain = load_qa_chain(model, chain_type="stuff", prompt=prompt)
        return chain
    except Exception as e:
        logger.error(f"Error getting question generation chain: {str(e)}")
        return f"Error getting question generation chain: {str(e)}"

def generate_top_questions():
    try:
        user_vector_store_path = session.get('user_vector_store_path')
        embeddings_model = OpenAIEmbeddings(openai_api_key=api_key2)
        new_db = FAISS.load_local(user_vector_store_path, embeddings_model, allow_dangerous_deserialization=True)
        docs = new_db.similarity_search("Generate the top 5 questions", k=10)
        chain = get_question_generation_chain()
        response = chain({"input_documents": docs, "context": ""}, return_only_outputs=True)
        return response["output_text"]
    except Exception as e:
        logger.error(f"Error generating top questions: {str(e)}")
        return f"Error generating top questions: {str(e)}"

def user_input(user_question):
    try:
        user_vector_store_path = session.get('user_vector_store_path')
        embeddings_model = OpenAIEmbeddings(openai_api_key=api_key2)
        new_db = FAISS.load_local(user_vector_store_path, embeddings_model, allow_dangerous_deserialization=True)
        docs = new_db.similarity_search(user_question, k=10)
        chain = get_conversation_chain()

        with get_openai_callback() as cb:
            response = chain({"input_documents": docs, "question": user_question}, return_only_outputs=True)
            cost = cb.total_cost
            tokens_used = cb.total_tokens
            prompt_tokens = cb.prompt_tokens
            completion_tokens = cb.completion_tokens

        logger.info(f"Tokens Used: {tokens_used}")
        logger.info(f"Prompt Tokens: {prompt_tokens}")
        logger.info(f"Completion Tokens: {completion_tokens}")
        logger.info(f"Successful Requests: {cb.successful_requests}")
        logger.info(f"Total Cost (USD): ${cost:.8f}")

        chat_history.append({"role": "user", "content": user_question})
        chat_history.append({"role": "assistant", "content": response["output_text"]})

        logger.info(f"Cost for generating response: ${cost:.4f}")
        
        user_id = session.get('user_id', 'unknown')
        
        # Return necessary details for logging
        return response["output_text"], tokens_used, prompt_tokens, completion_tokens, cb.successful_requests, cost, docs

    except Exception as e:
        logger.error(f"Error processing user input: {str(e)}")
        return f"Error processing user input: {str(e)}", None, None, None, None, None, None

def generateNewAnswer():
    try:
        for i in range(len(chat_history) - 1, -1, -1):
            if chat_history[i]["role"] == "user":
                last_user_question = chat_history[i]["content"]
                logger.info(f"Last user question extracted: {last_user_question}")
                break
        else:
            logger.warning("No previous user question found in chat history.")
            return "No previous user question found in chat history."

        new_prompt = f"""
        The user is not satisfied with the previous answer. Please provide a new response to the question: "{last_user_question}".
        Ensure the response is clear, concise, and addresses the question point by point or with a rephrased approach.
        """
        logger.info(f"New prompt generated for rephrased answer: {new_prompt}")

        with get_openai_callback() as cb:
            new_response = user_input(new_prompt)
            cost = cb.total_cost
            tokens_used = cb.total_tokens
            prompt_tokens = cb.prompt_tokens
            completion_tokens = cb.completion_tokens

        logger.info(f"Tokens Used: {tokens_used}")
        logger.info(f"Prompt Tokens: {prompt_tokens}")
        logger.info(f"Completion Tokens: {completion_tokens}")
        logger.info(f"Successful Requests: {cb.successful_requests}")
        logger.info(f"Total Cost (USD): ${cost:.8f}")

        logger.info(f"New response generated: {new_response}")
        logger.info(f"Cost for generating new response: ${cost:.4f}")

        return new_response
    except Exception as e:
        logger.error(f"Error generating new answer: {str(e)}")
        return f"Error generating new answer: {str(e)}"



def upload():
    try:
        auth_header = request.headers.get('Authorization')
        if not auth_header:
            return jsonify({"error": "Authorization header is missing."}), 401

        # Extract the token from the header
        token = auth_header.split(" ")[1]
          # Decode the JWT token
        try:
            decoded_token = jwt.decode(token, app.config['SECRET_KEY'], algorithms=["HS256"])
        except jwt.ExpiredSignatureError:
            return jsonify({"error": "Token has expired."}), 401
        except jwt.InvalidTokenError:
            return jsonify({"error": "Invalid token."}), 401
          # Extract user email from the decoded token
        user_email = decoded_token.get('email')
        if not user_email:
            return jsonify({"error": "User email not found in token."}), 401

        # Fetch user data from the database
        user_data = UserInfo.query.filter_by(Email=user_email).first()
        if not user_data:
            return jsonify({"error": "User data not found."}), 404

        uploaded_files = request.files.getlist("files")
        if not uploaded_files:
            return jsonify({"error": "No files uploaded"}), 400

        all_texts = ""
        for uploaded_file in uploaded_files:
            if uploaded_file.filename.lower().endswith(('.pdf', '.doc', '.docx', '.ppt', '.pptx', '.txt')):
                if uploaded_file.filename.lower().endswith('.pdf'):
                    all_texts += extract_text_from_pdf(uploaded_file.read())
                elif uploaded_file.filename.lower().endswith(('.doc', '.docx')):
                    all_texts += extract_text_from_docx(uploaded_file)
                elif uploaded_file.filename.lower().endswith(('.ppt', '.pptx')):
                    all_texts += extract_text_from_pptx(uploaded_file)
                elif uploaded_file.filename.lower().endswith('.txt'):
                    all_texts += uploaded_file.read().decode("utf-8")
            else:
                return jsonify({"error": "Unsupported file format"}), 400

        user_id = str(uuid.uuid4())
        session['user_id'] = user_id

        user_vector_store_path = f"C:/flask/persistent_storage/faiss_index_{user_id}"
        old_vector_store_path = session.get('user_vector_store_path')
        if old_vector_store_path:
            if os.path.exists(old_vector_store_path):
                shutil.rmtree(old_vector_store_path)

        session['user_vector_store_path'] = user_vector_store_path
        text_chunks = get_text_chunks(all_texts)
        get_vector_store(text_chunks, user_vector_store_path)

        # Log user activity including file upload
        log_activities(
            user_id=user_data.Pid,
            email=user_email,
            activity_type='File Upload',
            activity_content='Files uploaded successfully',
            response_content='Files processed and stored',
            prompt_tokens=None,
            completion_tokens=None,
            successful_requests=None,
            total_cost=None,
            # db=db,
            error_message=None,
            app_name='pdf'
        )

        return jsonify({"message": "Files uploaded successfully!"}), 200

    except Exception as e:
        logger.error(f"Error in upload route: {str(e)}")
        log_activities(
            user_id=user_data.Pid if user_data else None,
            email=user_email,
            app_name='pdf',
            activity_type='File Upload',
            activity_content=None,
            error_message=str(e)
        )
        return jsonify({"error": "An error occurred while processing your request."}), 500



def ask_question():
    try:
        auth_header = request.headers.get('Authorization')
        if not auth_header:
            return jsonify({"error": "Authorization header is missing."}), 401

        # Extract the token from the header
        token = auth_header.split(" ")[1]
          # Decode the JWT token
        try:
            decoded_token = jwt.decode(token, app.config['SECRET_KEY'], algorithms=["HS256"])
        except jwt.ExpiredSignatureError:
            return jsonify({"error": "Token has expired."}), 401
        except jwt.InvalidTokenError:
            return jsonify({"error": "Invalid token."}), 401
          # Extract user email from the decoded token
        user_email = decoded_token.get('email')
        if not user_email:
            return jsonify({"error": "User email not found in token."}), 401

        # Fetch user data from the database
        user_data = UserInfo.query.filter_by(Email=user_email).first()
        if not user_data:
            return jsonify({"error": "User data not found."}), 404

        data = request.json
        user_question = data.get('question')
        if not user_question:
            return jsonify({"error": "Question is required."}), 400

        user_vector_store_path = session.get('user_vector_store_path')
        if not user_vector_store_path or not os.path.exists(user_vector_store_path):
            return jsonify({"error": "Please upload files before asking a question."}), 400

        response_content, tokens_used, prompt_tokens, completion_tokens, successful_requests, cost, docs = user_input(user_question)
        
        # Log user activity including user input
        log_activities(
            user_data.Pid, user_email,
            app_name='pdf',
            activity_type='question',
            activity_content=user_question,
            response_content=response_content,
            total_cost=cost,
            successful_requests=successful_requests,
            completion_tokens=completion_tokens,
            prompt_tokens=prompt_tokens
        )

        return jsonify({"response": response_content}), 200

    except Exception as e:
        logger.error(f"Error in ask_question route: {str(e)}")
        log_activities(
            user_id=None, email=user_email,
            activity_type='error',
            app_name='ask_question',
            error_message=str(e)
        )
        return jsonify({"error": "An error occurred while processing your request."}), 500


def generate_top_questions_route():
    try:
        auth_header = request.headers.get('Authorization')
        if not auth_header:
            return jsonify({"error": "Authorization header is missing."}), 401

        # Extract the token from the header
        token = auth_header.split(" ")[1]
          # Decode the JWT token
        try:
            decoded_token = jwt.decode(token, app.config['SECRET_KEY'], algorithms=["HS256"])
        except jwt.ExpiredSignatureError:
            return jsonify({"error": "Token has expired."}), 401
        except jwt.InvalidTokenError:
            return jsonify({"error": "Invalid token."}), 401
          # Extract user email from the decoded token
        user_email = decoded_token.get('email')
        if not user_email:
            return jsonify({"error": "User email not found in token."}), 401

        # Fetch user data from the database
        user_data = UserInfo.query.filter_by(Email=user_email).first()
        if not user_data:
            return jsonify({"error": "User data not found."}), 404

        # Ensure user has uploaded files
        user_vector_store_path = session.get('user_vector_store_path')
        if not user_vector_store_path or not os.path.exists(user_vector_store_path):
            return jsonify({"error": "Please upload files before generating questions."}), 400

        response_content = generate_top_questions()
        
        # Log user activity including top questions generation
        log_activities(
            user_id=user_data.Pid,
            email=user_email,
            activity_type='Generate Top Questions',
            activity_content='Top questions generated successfully',
            response_content=response_content,
            app_name='pdf'
        )

        return jsonify({"top_questions": response_content}), 200

    except Exception as e:
        logger.error(f"Error generating top questions: {str(e)}")
        log_activities(
            user_id=user_data.Pid if user_data else None,
            email=user_email,
            activity_type='error',
            app_name='generate_top_questions',
            error_message=str(e)
        )
        return jsonify({"error": "An error occurred while generating top questions."}), 500



def dislike_feedback():
    try:
        auth_header = request.headers.get('Authorization')
        if not auth_header:
            return jsonify({"error": "Authorization header is missing."}), 401

        # Extract the token from the header
        token = auth_header.split(" ")[1]
          # Decode the JWT token
        try:
            decoded_token = jwt.decode(token, app.config['SECRET_KEY'], algorithms=["HS256"])
        except jwt.ExpiredSignatureError:
            return jsonify({"error": "Token has expired."}), 401
        except jwt.InvalidTokenError:
            return jsonify({"error": "Invalid token."}), 401
          # Extract user email from the decoded token
        user_email = decoded_token.get('email')
        if not user_email:
            return jsonify({"error": "User email not found in token."}), 401

        # Fetch user data from the database
        user_data = UserInfo.query.filter_by(Email=user_email).first()
        if not user_data:
            return jsonify({"error": "User data not found."}), 404

        # Get feedback data from request
        feedback_data = request.json
        chosen_option = feedback_data.get('option')
        
        # Check if the chosen option is valid
        options = {
            "1": "The response was inaccurate.",
            "2": "The response was unclear.",
            "3": "The response was too long.",
            "4": "The response was irrelevant."
        }

        if chosen_option not in options:
            return jsonify({"error": "Invalid option selected."}), 400

        # Get the full feedback message
        feedback_message = options[chosen_option]

        # Get the last user question and answer from chat history
        last_user_question = None
        last_user_answer = None

        for entry in reversed(chat_history):
            if entry["role"] == "user":
                last_user_question = entry["content"]
                break

        for entry in reversed(chat_history):
            if entry["role"] == "assistant":
                last_user_answer = entry["content"]
                break

        # Log the user feedback option
        log_activities(
            user_id=user_data.Pid,
            email=user_email,
            activity_type="Dislike Feedback",
            activity_content=feedback_message,
            response_content=None,
            app_name='pdf',
            error_message=None
        )

        logger.info("User feedback option logged")

        # Return response including the previous question, answer, and feedback
        return jsonify({
            "message": "Feedback recorded",
            "chosen_option": feedback_message,
            "previous_question": last_user_question,
            "previous_answer": last_user_answer
        }), 200

    except Exception as e:
        logger.error(f"Error handling dislike feedback: {str(e)}")
        log_activities(
            user_id=user_data.Pid if user_data else None,
            email=user_email if user_email else "Unknown",
            activity_type='error',
            app_name='dislike_feedback',
            error_message=str(e)
        )
        return jsonify({"error": "An error occurred while processing your request."}), 500



def get_chat_history():
    logger.info("Fetching chat history")
    return jsonify(chat_history)


def submit_feedback():
    try:
        auth_header = request.headers.get('Authorization')
        if not auth_header:
            return jsonify({"error": "Authorization header is missing."}), 401

        # Extract the token from the header
        token = auth_header.split(" ")[1]
          # Decode the JWT token
        try:
            decoded_token = jwt.decode(token, app.config['SECRET_KEY'], algorithms=["HS256"])
        except jwt.ExpiredSignatureError:
            return jsonify({"error": "Token has expired."}), 401
        except jwt.InvalidTokenError:
            return jsonify({"error": "Invalid token."}), 401
          # Extract user email from the decoded token
        user_email = decoded_token.get('email')
        if not user_email:
            return jsonify({"error": "User email not found in token."}), 401

        # Fetch user data from the database
        user_data = UserInfo.query.filter_by(Email=user_email).first()
        if not user_data:
            return jsonify({"error": "User data not found."}), 404

        # Get feedback data from request
        feedback_data = request.json
        is_liked = feedback_data.get('liked')

        if is_liked is None:
            return jsonify({"error": "Feedback (liked) is required."}), 400

        feedback_message = "Thank you for your feedback! We're glad you liked it." if is_liked else "Thank you for your feedback! We'll try to improve."

        # Log the user feedback
        log_activities(
            user_id=user_data.Pid,
            email=user_email,
            activity_type="Like Feedback" if is_liked else "Dislike Feedback",
            activity_content="User liked the response" if is_liked else "User disliked the response",
            response_content=None,
            app_name='pdf',
            error_message=None
        )

        logger.info("User feedback logged")

        # Return response
        return jsonify({"message": feedback_message}), 200

    except Exception as e:
        logger.error(f"Error handling feedback: {str(e)}")
        log_activities(
            user_id=user_data.Pid if user_data else None,
            email=user_email if user_email else "Unknown",
            activity_type='error',
            app_name='like_feedback',
            error_message=str(e)
        )
        return jsonify({"error": "An error occurred while processing your request."}), 500


# @app.route('/user/signin', methods=['POST'])
# def signin():
#     data = request.json
#     response, status_code = signin_user(data, app.secret_key)
#     return jsonify(response), status_code

# @app.route('/user/register', methods=['POST'])
# def register():
#     data = request.json
#     response, status_code = register_user(data)
#     return jsonify(response), status_code

# @app.route('/user/forgot', methods=['POST'])
# def forgot_password_route():
#     data = request.json
#     response, status_code = forgot_password(data)
#     return jsonify(response), status_code


if __name__ == "__main__":
    app.run()
